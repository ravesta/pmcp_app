"""
Mengubah berkas Excel "Project List" menjadi JSON datar yang dibaca skrip impor data actual.

Tujuan: memisahkan pembacaan Excel dari aturan pemetaan ke database. Berkas ini HANYA menyalin isi sel apa adanya
(tanggal menjadi YYYY-MM-DD, sel kosong atau "-" menjadi null); semua aturan bisnis ada di `import-actual-data.mjs`.
Pemakaian: python scripts/convert-project-list.py <berkas.xlsx> <keluaran.json>
Masukan: lembar "Daftar Project" (data mulai baris 3) dan lembar "PIC".
Efek samping: menulis berkas JSON keluaran; tidak menyentuh database.
Rujukan: lampiran BA "Project List_2026.xlsx", Technical Design (impor data actual).
"""
import datetime
import json
import sys

import openpyxl

# Urutan kolom lembar "Daftar Project" (indeks mulai 0) menjadi nama isian JSON.
COLUMNS = [
    'no', 'epic', 'name', 'ba', 'qa', 'requestDate',
    'baEstStart', 'baEstEnd', 'baActualStart', 'baActualEnd',
    'devEstStart', 'devEstEnd', 'devActualStart', 'devActualEnd',
    'sitEstStart', 'sitEstEnd', 'sitActualStart', 'sitActualEnd',
    'uatEstStart', 'uatEstEnd', 'uatActualStart', 'uatActualEnd',
    'uatSignOff', 'goLive', 'status', 'remarks1', 'remarks2', 'remarks3',
]


def clean(value):
    """Mengubah nilai sel menjadi string bersih, tanggal ISO, atau None. Masukan: nilai sel. Keluaran: nilai bersih."""
    if value is None:
        return None
    if isinstance(value, datetime.datetime):
        return value.date().isoformat()
    if isinstance(value, datetime.date):
        return value.isoformat()
    if isinstance(value, str):
        text = value.strip()
        return None if text in ('', '-') else text
    return value


def main():
    """Membaca Excel dan menulis JSON. Masukan: argumen baris perintah. Keluaran: tidak ada. Efek samping: menulis berkas."""
    source, target = sys.argv[1], sys.argv[2]
    book = openpyxl.load_workbook(source, data_only=True)

    pics = []
    for row in book['PIC'].iter_rows(min_row=3, values_only=True):
        name, stage = clean(row[1]), clean(row[2])
        if name:
            pics.append({'name': name, 'stage': stage})

    rows = []
    for row in book['Daftar Project'].iter_rows(min_row=3, values_only=True):
        if not any(cell is not None for cell in row):
            continue
        padded = list(row) + [None] * (len(COLUMNS) - len(row))
        rows.append({name: clean(padded[index]) for index, name in enumerate(COLUMNS)})

    with open(target, 'w', encoding='utf-8') as handle:
        json.dump({'source': source.replace('\\', '/').split('/')[-1], 'pics': pics, 'projects': rows}, handle, ensure_ascii=False, indent=1)
    print(f'{len(pics)} PIC dan {len(rows)} project ditulis ke {target}')


main()
